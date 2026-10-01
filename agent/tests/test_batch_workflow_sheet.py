"""The packaged tool's own batch sheet, and the board one row becomes.

`excel_mau_batch.xlsx` ships with the tool. Read rather than guessed (18/09):
sheet **`Workflow Batch`**, header **`image_prompt | video_prompt | Lời Thoại`**,
and every cell holds SEVERAL LINES — line *i* of the three columns is the same
scene. So a row is a video of N scenes, not a scene.

That is the shape `analyze_video` already writes (`branch_texts`, one line per
scene) and the shape `fan_out` already splits. This importer therefore adds no
third notion of "how a row becomes scenes"; it fills the same three fields.

The one thing it refuses to do is repair a misaligned row: padding would put a
line of dialogue under the wrong picture, and dropping one would lose a scene
without saying so.
"""
from __future__ import annotations

import io
from pathlib import Path

import pytest

from flowboard.db import get_session
from flowboard.db.models import Node

PACKAGED = Path("D:/TOOL_VIDEO/TOOL/excel_mau_batch.xlsx")


def _sheet(rows: list[tuple], *, title: str = "Workflow Batch") -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


HEADER = ("image_prompt", "video_prompt", "Lời Thoại")


def _post(client, payload: bytes, name: str = "batch.xlsx"):
    return client.post(
        "/api/batch/workflow",
        files={
            "file": (
                name,
                payload,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


# ── reading the sheet ─────────────────────────────────────────────────


def test_the_packaged_sample_parses(client):
    """The file the tool ships. If this stops parsing, the format moved."""
    if not PACKAGED.is_file():
        pytest.skip("the packaged sample sheet is not on this machine")
    resp = _post(client, PACKAGED.read_bytes(), PACKAGED.name)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["sheet"] == "Workflow Batch"
    assert body["rows"], body
    first = body["rows"][0]
    # Two scenes per row in the shipped sample, aligned across three columns.
    assert len(first["imagePrompts"]) == len(first["videoPrompts"]) == 2
    assert len(first["voicePrompts"]) == 2
    assert first["misaligned"] is False


def test_a_row_is_a_video_of_several_scenes(client):
    payload = _sheet([
        HEADER,
        ("cảnh 1 ảnh\ncảnh 2 ảnh", "cảnh 1 video\ncảnh 2 video", "thoại 1\nthoại 2"),
    ])
    row = _post(client, payload).json()["rows"][0]
    assert row["imagePrompts"] == ["cảnh 1 ảnh", "cảnh 2 ảnh"]
    assert row["voicePrompts"] == ["thoại 1", "thoại 2"]


def test_a_misaligned_row_is_reported_not_repaired(client):
    """Padding would put a line of dialogue under the wrong picture; dropping
    one would lose a scene in silence. Both are worse than saying so."""
    payload = _sheet([
        HEADER,
        ("ảnh 1\nảnh 2\nảnh 3", "video 1\nvideo 2", "thoại 1"),
    ])
    row = _post(client, payload).json()["rows"][0]
    assert row["misaligned"] is True
    assert len(row["imagePrompts"]) == 3 and len(row["videoPrompts"]) == 2


def test_blank_rows_are_counted_not_dropped_silently(client):
    payload = _sheet([HEADER, ("ảnh", "video", "thoại"), (None, None, None)])
    body = _post(client, payload).json()
    assert len(body["rows"]) == 1
    assert body["skipped"] == 1


def test_a_sheet_that_is_not_the_packaged_format_says_so(client):
    payload = _sheet([("ghi chú", "ngày"), ("x", "y")], title="Sheet1")
    resp = _post(client, payload)
    assert resp.status_code == 400
    assert "Workflow Batch" in resp.json()["detail"]


def test_the_sheet_is_found_by_name_among_several(client):
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Ghi chú"
    book.active.append(["không phải cột nào cả"])
    sheet = book.create_sheet("Workflow Batch")
    for row in (HEADER, ("ảnh 1", "video 1", "thoại 1")):
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)

    body = _post(client, buffer.getvalue()).json()
    assert body["sheet"] == "Workflow Batch"
    assert body["rows"][0]["videoPrompts"] == ["video 1"]


# ── the board a row becomes ───────────────────────────────────────────


def _board(client) -> int:
    return client.post("/api/boards", json={"name": "Excel"}).json()["id"]


def test_a_row_becomes_a_board_carrying_its_three_lists(client):
    board_id = _board(client)
    resp = client.post("/api/batch/workflow/board", json={
        "board_id": board_id,
        "row": {
            "index": 1,
            "imagePrompts": ["ảnh 1", "ảnh 2"],
            "videoPrompts": ["video 1", "video 2"],
            "voicePrompts": ["thoại 1", "thoại 2"],
        },
    })
    assert resp.status_code == 200, resp.text
    built = resp.json()
    assert built["scenes"] == 2

    with get_session() as session:
        node = session.get(Node, built["promptNodeId"])
        data = node.data or {}
    # One line per scene, in the three fields the branch sockets read.
    assert data["imagePrompt"].splitlines() == ["ảnh 1", "ảnh 2"]
    assert data["videoPrompt"].splitlines() == ["video 1", "video 2"]
    assert data["voicePrompt"].splitlines() == ["thoại 1", "thoại 2"]
    # And `prompt` too, because that is what the fan-out counts lines of.
    assert data["prompt"].splitlines() == ["ảnh 1", "ảnh 2"]


def test_importing_prices_the_board_but_dispatches_nothing(client):
    """Unlike an empty archetype board, this one arrives with prompts — so it
    IS ready to spend, and the estimate has to say so before anyone presses
    Run. What importing must never do is dispatch: no Request row exists until
    the user runs the board."""
    from sqlmodel import select

    from flowboard.db.models import Request

    board_id = _board(client)
    client.post("/api/batch/workflow/board", json={
        "board_id": board_id,
        "row": {"index": 1, "imagePrompts": ["a"], "videoPrompts": ["b"], "voicePrompts": []},
    })
    estimate = client.get(f"/api/boards/{board_id}/estimate").json()
    assert estimate["billableJobs"] == 2, "one image + one video, both with text"
    assert estimate["notReadyJobs"] == 0

    with get_session() as session:
        assert session.exec(select(Request)).all() == []


def test_the_scene_count_follows_the_row(client):
    board_id = _board(client)
    built = client.post("/api/batch/workflow/board", json={
        "board_id": board_id,
        "row": {
            "index": 1,
            "imagePrompts": ["a", "b", "c"],
            "videoPrompts": ["a", "b", "c"],
            "voicePrompts": ["x", "y", "z"],
        },
    }).json()
    assert built["scenes"] == 3
    detail = client.get(f"/api/boards/{board_id}").json()
    assert sum(1 for n in detail["nodes"] if n["type"] == "video") == 3


def test_an_empty_row_is_refused(client):
    board_id = _board(client)
    resp = client.post("/api/batch/workflow/board", json={
        "board_id": board_id,
        "row": {"index": 1, "imagePrompts": [], "videoPrompts": [], "voicePrompts": []},
    })
    assert resp.status_code == 422


def test_a_missing_board_is_a_404(client):
    resp = client.post("/api/batch/workflow/board", json={
        "board_id": 9999,
        "row": {"index": 1, "imagePrompts": ["a"], "videoPrompts": ["a"], "voicePrompts": []},
    })
    assert resp.status_code == 404
